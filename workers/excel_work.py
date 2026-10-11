from openpyxl import load_workbook
import time
from datetime import datetime
from api.cnpja_api import CnpjaApi, CnpjaApiError
from cleaners.cnpj import CnpjCleaner
from parsers.cnpj.cnpja_parser import CnpjaParser


class ExcelProcessor:
    def __init__(self, file_path):
        self.file_path = file_path

    def process(self):
        workbook = load_workbook(self.file_path)
        worksheet = workbook.active

        for row in range(2, worksheet.max_row + 1):
            cnpj = worksheet.cell(row=row, column=1).value

            if not cnpj:
                continue

            self._process_row(worksheet, row, cnpj)

        workbook.save(self.file_path)

    def _process_row(self, worksheet, row, cnpj):
        cnpj = CnpjCleaner.normalize(cnpj)

        if not CnpjCleaner.is_valid(cnpj):
            worksheet.cell(row=row, column=2).value = "CNPJ inválido"
            print(f"O CNPJ {cnpj} é inválido ou contém um erro de digitação.")
            return

        try:
            time.sleep(15)
            data = CnpjaApi.get_company(cnpj)
            client = CnpjaParser.parse(data)

        except CnpjaApiError:
            worksheet.cell(row=row, column=2).value = "Erro na consulta"
            print(f"Não encontramos dados do CNPJ {cnpj}.")
            return

        self._write_client(worksheet, row, client)

    @staticmethod
    def _write_client(worksheet, row, client):
        worksheet.cell(row=row, column=1).value = client.cnpj
        worksheet.cell(row=row, column=2).value = client.razao
        worksheet.cell(row=row, column=3).value = client.rua
        worksheet.cell(row=row, column=4).value = client.num
        worksheet.cell(row=row, column=5).value = client.bairro
        worksheet.cell(row=row, column=6).value = client.municipio
        worksheet.cell(row=row, column=7).value = client.estado
        worksheet.cell(row=row, column=8).value = client.cep
        worksheet.cell(row=row, column=9).value = client.email
        worksheet.cell(row=row, column=10).value = client.fone
        print(f"Processando...\nÚltimo Cliente processado: {client.razao}\nÚltima atualização: {datetime.now()}")
